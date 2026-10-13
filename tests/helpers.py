"""Fixture-file builders shared by all test modules. Files are generated in code
so the repo needs no binary fixtures."""
from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas

LONG = "This is a digital page with plenty of real extractable text on it."


def make_pdf(path: Path, pages: list[str | None]) -> Path:
    """Each entry is one page. A str gives a digital page; None gives a page with
    no text layer, which the extractor should treat as scanned."""
    c = canvas.Canvas(str(path))
    for text in pages:
        if text:
            c.drawString(72, 720, text)
        else:
            c.rect(10, 10, 5, 5)  # draws a shape, but no text
        c.showPage()
    c.save()
    return path


def make_docx(path: Path, paragraphs: list[str] | None = None) -> Path:
    d = Document()
    if paragraphs is None:
        d.add_heading("Title here", 1)
        d.add_paragraph("Body paragraph.")
        t = d.add_table(rows=2, cols=2)
        for r, row in enumerate([["a", "b"], ["1", "2"]]):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
    else:
        for p in paragraphs:
            d.add_paragraph(p)
    d.save(path)
    return path
